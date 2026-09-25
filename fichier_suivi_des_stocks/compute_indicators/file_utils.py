import copy
import os
import re
from datetime import datetime
from pathlib import Path, PosixPath

import numpy as np
import openpyxl as pyxl
import pandas as pd
from openhexa.sdk import workspace
from openpyxl.worksheet.dimensions import ColumnDimension, RowDimension

from .fetch_pa_from_qat import extract_pa
import contextlib


def process_pa_files(
    fp_plan_approv: PosixPath,
    fp_map_prod: PosixPath,
    programme: str,
    date_report: str,
) -> pd.DataFrame:
    """Process and merge plan approval files with product mapping data.

    Args:
        fp_plan_approv (PosixPath): Path to the directory or file containing plan approval CSV files.
        fp_map_prod (PosixPath): Path to the Excel file containing product mapping data.
        programme (str): The sheet name in the Excel file to be used for product mapping.
        date_report (str): The date of the report in the format "YYYY-MM-DD".

    Returns:
        pd.DataFrame: A DataFrame containing the processed and merged data.
    """
    if (
        fp_plan_approv is None
        or fp_plan_approv
        == Path(workspace.files_path)
        / f"Fichier Suivi de Stock/data/{programme}/Plan d'Approvisionnement"
    ):
        conn = workspace.get_connection("qat")
        credentials = {"username": conn.username, "password": conn.password}  # type: ignore
        df_plan_approv = extract_pa(
            programme_id=programme, credentials=credentials, date_report=date_report
        )
        df_plan_approv["date de réception"] = df_plan_approv["date de réception"].apply(
            lambda d: datetime.strptime(d, "%Y-%m-%d")
        )

    else:
        # Liste pour accumuler les données
        data = []

        # Fonction pour traiter un fichier unique
        def _process_pa_file(fichier: Path) -> None:
            nonlocal data
            version = None
            with Path(fichier).open(encoding="utf-8") as file:
                for line in file.readlines():
                    if version is None and "version" in line.lower():
                        try:
                            version = line.split(":")[-1].replace('"', "").strip()
                            list_element = _process_pa_version(version)
                        except IndexError:
                            version = ""

                    if len(line.split(",", 16)) == 17:
                        data.append(
                            [col.replace('"', "").strip() for col in line.split(",", 16)]
                            + list_element
                        )

        if Path(fp_plan_approv).is_dir():
            for root, _, files in os.walk(fp_plan_approv):
                for file in files:
                    if file.endswith(".csv"):
                        _process_pa_file(Path(root) / file)
        else:
            _process_pa_file(fp_plan_approv)

        df_plan_approv = pd.DataFrame(
            data[1:], columns=data[0][:-2] + ["version_pa", "date_extraction_pa"]
        )

        # Bad index
        bad_index = df_plan_approv.loc[
            df_plan_approv["ID de produit QAT / Identifiant de produit (prévision)"]
            == "ID de produit QAT / Identifiant de produit (prévision)"
        ].index

        df_plan_approv = df_plan_approv.drop(index=bad_index)

        for col in ["ID de produit QAT / Identifiant de produit (prévision)", "ID de l`envoi QAT"]:
            with contextlib.suppress(Exception):
                df_plan_approv[col] = df_plan_approv[col].astype("Int64")

    for col in [
        "Coût unitaire de produit (USD)",
        "Coût du fret (USD)",
        "Quantité",
        "Coût total (USD)",
    ]:
        with contextlib.suppress(Exception):
            df_plan_approv[col] = df_plan_approv[col].astype(float)

    with contextlib.suppress(Exception):
        df_plan_approv["date de réception"] = df_plan_approv["date de réception"].apply(
            lambda date_str: datetime.strptime(date_str, "%d-%b-%Y")
        )

    # Nettoyage des espaces blancs dans les données
    df_plan_approv = df_plan_approv.apply(lambda x: x.str.strip() if x.dtype == "object" else x)

    # Charger le fichier de mappage des produits
    df_map_prod = pd.read_excel(fp_map_prod, sheet_name=programme)  # ou pd.read_csv selon le type
    df_map_prod.columns = (
        df_map_prod.columns.str.replace("Ã©", "é").str.replace("â", "").str.rstrip().str.lstrip()
    )

    df_map_prod = df_map_prod.rename(
        columns={
            "Code QAT": "ID de produit QAT / Identifiant de produit (prévision)",
            "Code standard national": "Standard product code",
            "Coût unitaire moyen (en dollar)": "cout_unitaire_moyen_qat",
            "Facteur de conversion QAT vers SAGE": "facteur_de_conversion_qat_sage",
            "Acronym": "acronym",
        },
    )

    df_map_prod = df_map_prod.drop_duplicates()

    df_plan_approv = df_plan_approv.merge(
        df_map_prod[
            [
                "ID de produit QAT / Identifiant de produit (prévision)",
                "Standard product code",
                "facteur_de_conversion_qat_sage",
                "acronym",
            ]
        ],
        on="ID de produit QAT / Identifiant de produit (prévision)",
        how="left",
    )

    df_plan_approv = df_plan_approv.rename(
        columns={
            "ID de produit QAT / Identifiant de produit (prévision)": "ID de produit QAT",
            "Produit (planification) / Produit (prévision)": "Produits",
            "État": "Status",
            "Quantité": "Quantite",
            "date de réception": "DATE",
            "Coût unitaire de produit (USD)": "Cout des Produits",
            "Coût du fret (USD)": "Couts du fret",
            "Coût total (USD)": "Couts totaux",
        },
    )

    df_plan_approv["cout_unitaire_moyen_qat"] = (
        (df_plan_approv["Cout des Produits"] / df_plan_approv["Quantite"])
        .replace([np.inf, -np.inf], np.nan)
        .fillna(0)
    )

    df_plan_approv["code_and_date_concate"] = df_plan_approv.apply(
        _get_code_and_date_concate, axis=1
    )

    return df_plan_approv


def process_etat_stock_npsp(
    df_etat_stock_npsp: pd.DataFrame, date_report: str, programme: str
) -> pd.DataFrame:
    """Traite le DataFrame de l'état des stocks NPSP en renommant les colonnes,
    en ajoutant la date du rapport et le programme, et en sélectionnant les colonnes pertinentes.

    Args:
        df_etat_stock_npsp (pd.DataFrame): Le DataFrame contenant les données de l'état des stocks NPSP.
        date_report (str): La date du rapport au format 'YYYY-MM-DD'.
        programme (str): Le programme associé aux données.

    Returns:
        pd.DataFrame: Le DataFrame traité avec les colonnes renommées, la date du rapport et le programme ajoutés.
    """  # noqa: D205, E501
    COLUMN_MAPPING = {  # noqa: N806
        r"Nouveau code": "code_produit",
        r"Nouvelle désignation": "designation",
        r"Contenance": "contenance",
        r"DMM": "dmm",
        r"Traceurs": "traceurs",
        r"MSD": "msd",
        r"Statut du stock": "statut_stock",
        r"Stock théorique BOUAKE": "stock_theorique_bke",
        r"Stock théorique ABIDJAN": "stock_theorique_abj",
        r"Stock théorique CENTRALE": "stock_theorique_central",
        r"Stock théorique fin": "stock_theorique_fin_mois",
        r"Nombre de jour de rupture": "nb_jour_rupture",
    }

    df_etat_stock_npsp = df_etat_stock_npsp.rename(
        columns=lambda col: next(
            (v for k, v in COLUMN_MAPPING.items() if re.search(k, col, re.I)), col
        ),
    )

    df_etat_stock_npsp["date_report"] = pd.to_datetime(date_report, format="%Y-%m-%d")

    cols = [
        col
        for col in [
            "date_report",
            "code_produit",
            "designation",
            "contenance",
            "dmm",
            "traceurs",
            "stock_theorique_bke",
            "stock_theorique_abj",
            "stock_theorique_central",
            "stock_theorique_fin_mois",
            "msd",
            "statut_stock",
        ]
        if col in df_etat_stock_npsp.columns
    ]
    df_etat_stock_npsp = df_etat_stock_npsp[cols]

    df_etat_stock_npsp["programme"] = programme

    return df_etat_stock_npsp


def _get_code_and_date_concate(row) -> str | float:
    try:
        if not pd.isna(row["Standard product code"]):
            return (
                str(int(row["Standard product code"]))
                + "_"
                + str(row["DATE"]).replace(" 00:00:00", "")
            )
        if pd.isna(row["Standard product code"]) and not pd.isna(row["DATE"]):
            return "_" + str(row["DATE"]).replace(" 00:00:00", "")
        return np.nan
    except Exception:
        try:
            return "_" + str(row["DATE"]).replace(" 00:00:00", "")
        except Exception:
            return np.nan


DATE_EXTRACT_PATTERN = re.compile(r"\((\w{3,9} \d{1,2} \d{4})\)")


def _process_pa_version(pa_version: str):
    """Extracts the version number and date from the plan approval version string.

    Args:
        pa_version (str): The version string of the plan approval.

    Returns:
        list: A list containing the version number and date as a datetime object.
    """
    number, date_obj = None, None

    # Utilisé pour extraire le nombre de début
    if num_match := re.match(r"^\d+", pa_version):
        number = int(num_match.group())

    # Extraction de la date de la version du fichier du plan d'approvisionnement
    if date_match := DATE_EXTRACT_PATTERN.findall(pa_version):
        date_str = date_match[-1]
        for fmt in ("%b %d %Y", "%B %d %Y", "%d-%b-%Y", "%d %b %Y"):
            try:
                date_obj = datetime.strptime(date_str, fmt)
                break
            except ValueError:
                continue

    return [number, date_obj]


def _find_sheet_by_keyword(wb: pyxl.Workbook, keyword: str) -> str | None:
    """Retrouve le nom exact d'une feuille contenant un mot-clé (insensible à la casse).

    Args:
        wb (pyxl.Workbook): Le classeur dans lequel rechercher la feuille.
        keyword (str): Le mot-clé à rechercher dans les noms de feuilles.

    Returns:
        str | None: Le nom exact de la feuille si trouvée, sinon None.
    """
    return next((name for name in wb.sheetnames if keyword.lower() in name.lower()), None)


def _copy_sheet(
    src_wb: pyxl.Workbook, src_sheet_name: str, dest_wb: pyxl.Workbook, new_sheet_name: str
) -> None:
    """Copie une feuille (valeurs, formules, styles, fusions, largeurs de colonnes)
    d'un classeur source vers un classeur destination.

    Args:
        src_wb (pyxl.Workbook): Le classeur source (ex: le template).
        src_sheet_name (str): Le nom de la feuille à copier depuis le classeur source.
        dest_wb (pyxl.Workbook): Le classeur destination (ex: le fichier programme).
        new_sheet_name (str): Le nom à donner à la feuille copiée dans le classeur destination.
    """  # noqa: D205
    src_ws = src_wb[src_sheet_name]
    dest_ws = dest_wb.create_sheet(title=new_sheet_name)

    for row in src_ws.iter_rows():
        for cell in row:
            new_cell = dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = copy.copy(cell.font)
                new_cell.border = copy.copy(cell.border)
                new_cell.fill = copy.copy(cell.fill)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy.copy(cell.protection)
                new_cell.alignment = copy.copy(cell.alignment)

    for merged_range in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(merged_range))

    for col_letter, dim in src_ws.column_dimensions.items():
        dest_ws.column_dimensions[col_letter] = ColumnDimension(
            dest_ws, index=col_letter, width=dim.width, hidden=dim.hidden
        )

    for row_idx, dim in src_ws.row_dimensions.items():
        dest_ws.row_dimensions[row_idx] = RowDimension(
            dest_ws, index=row_idx, height=dim.height, hidden=dim.hidden
        )

    dest_ws.freeze_panes = src_ws.freeze_panes


def _fix_etat_de_stock_headers(wb: pyxl.Workbook) -> None:
    """Renomme les en-têtes de la feuille "Etat de stock <programme>" pour qu'ils
    correspondent au template : "Code" -> "Nouveau code" et la 2e occurrence de
    "Désignation" -> "Nouvelle désignation". Ne fait rien si les en-têtes sont
    déjà conformes.

    Args:
        wb (pyxl.Workbook): Le classeur du fichier programme (modifié en place).
    """  # noqa: D205
    sheet_name = _find_sheet_by_keyword(wb, "Etat de stock")
    if sheet_name is None:
        return

    ws = wb[sheet_name]

    header_row = next(
        (row for row in range(1, 10) if ws.cell(row=row, column=1).value == "Programme"),
        None,
    )
    if header_row is None:
        return

    designation_seen = False
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row, column=col)
        if cell.value == "Désignation":
            if not designation_seen:
                designation_seen = True
            else:
                cell.value = "Nouvelle désignation"
        elif cell.value == "Code":
            cell.value = "Nouveau code"


def _remove_last_column(wb: pyxl.Workbook) -> None:
    """Vide la cellule d'en-tête vide qui suit immédiatement la dernière colonne
    d'en-tête (avec contenu) sur la feuille "Etat de stock <programme>",
    après adaptation des en-têtes. Cette colonne est repérée par rapport à
    la ligne d'en-tête (celle contenant "Programme" en première colonne).

    Contrairement à une suppression de colonne complète (qui décale toutes
    les lignes et peut faire remonter par erreur de vraies données situées
    plus loin sur d'autres lignes), seule la cellule de la ligne d'en-tête
    est vidée (valeur et mise en forme), sans aucun décalage des autres
    lignes ni des autres colonnes.

    """  # noqa: D205
    sheet_name = _find_sheet_by_keyword(wb, "Etat de stock")
    if sheet_name is None:
        return

    ws = wb[sheet_name]

    # La ligne d'en-tête n'est pas à un numéro de ligne fixe : selon les
    # programmes, le nombre de lignes de titre au-dessus varie. On la
    # détecte dynamiquement en cherchant la ligne où la colonne A vaut
    # "Programme", plutôt que de coder en dur un numéro de ligne.
    header_row = next(
        (row for row in range(1, 10) if ws.cell(row=row, column=1).value == "Programme"),
        None,
    )
    if header_row is None:
        return

    last_header_col = max(
        (
            col
            for col in range(1, ws.max_column + 1)
            if ws.cell(row=header_row, column=col).value is not None
        ),
        default=None,
    )
    if last_header_col is None:
        return

    col_to_clear = last_header_col + 1
    if col_to_clear > ws.max_column:
        return

    cell = ws.cell(row=header_row, column=col_to_clear)
    cell.value = None
    cell.fill = pyxl.styles.PatternFill(fill_type=None)


def _adapt_ppi_sheet(template_wb: pyxl.Workbook, programme_wb: pyxl.Workbook) -> None:
    """Quand la feuille "PPI" existe déjà dans le fichier programme, adapte ses
    colonnes pour qu'elles correspondent exactement à celles du template :
    - renomme et réordonne les colonnes reconnues (Code Produit, Nom Produit,
      Unite, Numéro Lot, Date Peremption, Quantité) selon des alias courants
      (ex: ARTICLE -> Code Produit, DLC -> Date Peremption, etc.)
    - supprime les colonnes du programme qui n'ont pas d'équivalent dans le
      template (ex: DATE PIECE, N°PIECE, SOUS_LOT, EMPLACEMENT)
    - conserve toutes les lignes de données existantes, réorganisées selon
      l'ordre des colonnes du template
    Si la feuille "PPI" n'existe pas dans le fichier programme, cette fonction
    ne fait rien (c'est _ensure_ppi_sheet qui gère la copie depuis le template
    dans ce cas).

    Args:
        template_wb (pyxl.Workbook): Le classeur du template de référence.
        programme_wb (pyxl.Workbook): Le classeur du fichier programme (modifié en place).
    """  # noqa: D205
    sheet_name = _find_sheet_by_keyword(programme_wb, "PPI")
    if sheet_name is None:
        return

    template_sheet_name = _find_sheet_by_keyword(template_wb, "PPI")
    if template_sheet_name is None:
        return

    ws = programme_wb[sheet_name]
    tpl_ws = template_wb[template_sheet_name]

    tpl_header_row = next(
        (row for row in range(1, 10) if tpl_ws.cell(row=row, column=1).value == "Code Produit"),
        None,
    )
    if tpl_header_row is None:
        return

    template_columns = []
    col = 1
    while True:
        val = tpl_ws.cell(row=tpl_header_row, column=col).value
        if val is None:
            break
        template_columns.append(val)
        col += 1

    COLUMN_ALIASES = {  # noqa: N806
        # Chaque programme (PNLP, PNLT, PNLS...) exporte sa feuille PPI avec
        # des noms de colonnes différents selon l'outil/la personne qui l'a
        # produite (ex: "ARTICLE" au lieu de "Code Produit", "DLC" au lieu
        # de "Date Peremption"). Ce dictionnaire centralise les variantes
        # connues à ce jour. S'il apparaît un nouveau programme utilisant
        # encore un autre nom pour une colonne du template, c'est ici qu'il
        # faut ajouter le nouvel alias.
        "Code Produit": ["ARTICLE", "CODE PRODUIT", "CODE_PRODUIT"],
        "Nom Produit": ["DESIGNATION", "NOM PRODUIT", "NOM_PRODUIT"],
        "Unite": ["UNITE", "UNITÉ"],
        "Numéro Lot": ["LOT", "NUMERO LOT", "NUMÉRO LOT", "N° LOT"],
        "Date Peremption": [
            "DLC",
            "DATE PEREMPTION",
            "DATE DE PEREMPTION",
            "DATE LIMITE DE CONSOMMATION",
        ],
        "Quantité": ["QTE_STOCKAGE", "QUANTITE", "QTE"],
    }

    def _normalize(value: object) -> str:
        return str(value).strip().upper() if value is not None else ""

    prog_header_row = None
    all_aliases = {alias for aliases in COLUMN_ALIASES.values() for alias in aliases}
    for row in range(1, 10):
        values = [_normalize(ws.cell(row=row, column=c).value) for c in range(1, ws.max_column + 1)]
        matches = sum(1 for v in values if v in all_aliases)
        if matches >= 3:
            prog_header_row = row
            break
    if prog_header_row is None:
        return

    col_to_template = {}
    for c in range(1, ws.max_column + 1):
        val = _normalize(ws.cell(row=prog_header_row, column=c).value)
        for tpl_col, aliases in COLUMN_ALIASES.items():
            if val in aliases or val == tpl_col.upper():
                col_to_template[c] = tpl_col
                break

    if not col_to_template:
        return

    data_rows = []
    for r in range(prog_header_row + 1, ws.max_row + 1):
        row_vals = {col_to_template[c]: ws.cell(row=r, column=c).value for c in col_to_template}
        if any(v is not None for v in row_vals.values()):
            data_rows.append(row_vals)

    title = ws.cell(row=1, column=1).value

    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))

    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell.number_format = "General"

    if title:
        ws.cell(row=1, column=1, value=title)

    for idx, col_name in enumerate(template_columns, start=1):
        cell = ws.cell(row=tpl_header_row, column=idx, value=col_name)
        tpl_cell = tpl_ws.cell(row=tpl_header_row, column=idx)
        if tpl_cell.has_style:
            cell.font = copy.copy(tpl_cell.font)
            cell.fill = copy.copy(tpl_cell.fill)
            cell.border = copy.copy(tpl_cell.border)
            cell.alignment = copy.copy(tpl_cell.alignment)
            cell.number_format = tpl_cell.number_format

    for r_offset, row_vals in enumerate(data_rows, start=1):
        for c_idx, col_name in enumerate(template_columns, start=1):
            data_cell = ws.cell(
                row=tpl_header_row + r_offset, column=c_idx, value=row_vals.get(col_name)
            )
            if col_name == "Date Peremption":
                data_cell.number_format = "DD/MM/YYYY"

    # Supprime physiquement les colonnes en trop (sans en-tête / sans équivalent
    # dans le template), au-delà du nombre de colonnes du template.
    #
    # NOTE DE CONCEPTION : contrairement à _remove_last_column() (sur la
    # feuille "Etat de stock"), on utilise ici ws.delete_cols() sans risque,
    # car la feuille PPI est intégralement reconstruite juste au-dessus
    # (toutes les valeurs sont d'abord vidées puis réécrites depuis
    # data_rows). Il n'y a donc aucune donnée "existante ailleurs" qui
    # pourrait remonter par erreur à cette position après suppression.
    nb_template_cols = len(template_columns)
    if ws.max_column > nb_template_cols:
        ws.delete_cols(nb_template_cols + 1, ws.max_column - nb_template_cols)

    # Corrige la référence de l'AutoFilter (les flèches de filtre Excel),
    # qui n'est pas automatiquement réduite par delete_cols et continuerait
    # sinon d'afficher des flèches de filtre sur les colonnes supprimées.
    # (Bug/limitation connue d'openpyxl : delete_cols() décale les cellules
    # mais ne met pas à jour ws.auto_filter.ref, qui reste sur son ancienne
    # plage — ex: toujours "A3:J59" alors qu'il ne reste que 6 colonnes.)
    last_col_letter = pyxl.utils.get_column_letter(nb_template_cols)
    last_row = max(ws.max_row, tpl_header_row)
    ws.auto_filter.ref = f"A{tpl_header_row}:{last_col_letter}{last_row}"


def _ensure_ppi_sheet(template_wb: pyxl.Workbook, programme_wb: pyxl.Workbook) -> None:
    """Ajoute la feuille "PPI" (copiée depuis le template) au fichier programme si
    elle n'y existe pas déjà.

    Args:
        template_wb (pyxl.Workbook): Le classeur du template de référence.
        programme_wb (pyxl.Workbook): Le classeur du fichier programme (modifié en place).
    """  # noqa: D205
    if _find_sheet_by_keyword(programme_wb, "PPI") is not None:
        return

    template_ppi_name = _find_sheet_by_keyword(template_wb, "PPI")
    if template_ppi_name is None:
        return

    _copy_sheet(template_wb, template_ppi_name, programme_wb, new_sheet_name="PPI")


def adapt_programme_file_to_template(
    fp_etat_mensuel: PosixPath, template_path: PosixPath
) -> PosixPath:
    """Adapte le fichier "Etat Mensuel" d'un programme au format du template de référence :
    - renomme les en-têtes de la feuille "Etat de stock" si nécessaire
      (Code -> Nouveau code, Désignation -> Nouvelle désignation)
    - vide la cellule d'en-tête vide juste après le dernier champ renseigné
      sur la feuille "Etat de stock" (sans décaler aucune donnée)
    - si la feuille "PPI" existe déjà, adapte ses colonnes au format du
      template (renommage, réordonnancement, suppression des colonnes en trop)
    - si la feuille "PPI" n'existe pas, la copie depuis le template
    Le fichier original (fp_etat_mensuel) n'est jamais modifié : le résultat
    est enregistré dans un nouveau fichier, dans le même dossier, avec le
    suffixe "_final" ajouté au nom.

    NOTE DE CONCEPTION : on écrit toujours dans un nouveau fichier plutôt que
    d'écraser fp_etat_mensuel, pour deux raisons : (1) si l'adaptation
    échoue ou produit un résultat inattendu, le fichier programme source
    (souvent la seule copie disponible) reste intact et récupérable ;
    (2) ça facilite le débogage en gardant les deux versions (avant/après)
    consultables côte à côte.

    Args:
        fp_etat_mensuel (PosixPath): Chemin du fichier "Etat Mensuel" du programme à adapter.
        template_path (PosixPath): Chemin du fichier template de référence.

    Returns:
        PosixPath: Le chemin du nouveau fichier programme corrigé.
    """  # noqa: D205
    fp_etat_mensuel = Path(fp_etat_mensuel)  # type: ignore

    template_wb = pyxl.load_workbook(template_path)
    programme_wb = pyxl.load_workbook(fp_etat_mensuel)

    _fix_etat_de_stock_headers(programme_wb)
    _remove_last_column(programme_wb)
    _adapt_ppi_sheet(template_wb, programme_wb)
    _ensure_ppi_sheet(template_wb, programme_wb)

    output_path = fp_etat_mensuel.with_name(f"{fp_etat_mensuel.stem}_final{fp_etat_mensuel.suffix}")
    programme_wb.save(output_path)

    return output_path
